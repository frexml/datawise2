"""Source-to-Target mapping register export (Excel).

Follows the DataWise Platform Development Guide's generate_s2t_register
pattern: one workbook per job, reading from the already-computed Lineage
table (the same rows GET /api/jobs/{id}/lineage returns).
"""
import io

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from dsxlineage.db import models

_HEADERS = [
    "Source Table", "Source Field", "Source Link",
    "Target Table", "Target Field", "Target Link",
    "Transformation Logic", "Transformation Type", "Cardinality", "Total Hops",
]

_HEADER_FILL = PatternFill(start_color="1A1C2B", end_color="1A1C2B", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def generate_s2t_register(job_id: int, db: Session) -> io.BytesIO:
    job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not job:
        raise ValueError(f"Job {job_id} not found")

    rows = db.query(models.Lineage).filter(models.Lineage.job_id == job_id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (job.filename or f"job_{job_id}")[:31]  # Excel sheet name limit

    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for r, row in enumerate(rows, start=2):
        ws.cell(r, 1, row.source_table)
        ws.cell(r, 2, row.source_field)
        ws.cell(r, 3, row.source_link)
        ws.cell(r, 4, row.target_table)
        ws.cell(r, 5, row.target_field)
        ws.cell(r, 6, row.target_link)
        ws.cell(r, 7, row.transformation_logic)
        ws.cell(r, 8, row.transformation_type)
        ws.cell(r, 9, row.cardinality)
        ws.cell(r, 10, row.total_hops)

    for col in range(1, len(_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
