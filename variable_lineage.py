import pandas as pd
import requests
import sys
import json
import re
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def extract_columns_from_link(link_properties):
    """Extract column mappings from link properties"""
    columns = []
    
    if not link_properties or not isinstance(link_properties, dict):
        return columns
    
    # Look for schema/column information in link properties
    if 'Columns' in link_properties:
        cols = link_properties['Columns']
        if isinstance(cols, list):
            for col in cols:
                if isinstance(col, dict):
                    columns.append({
                        'name': col.get('Name', ''),
                        'type': col.get('SqlType', ''),
                        'derivation': col.get('Derivation', '')
                    })
    
    return columns

def extract_stage_columns(stage_properties):
    """Extract columns from stage properties (transformations, derivations)"""
    columns = {}
    
    if not stage_properties or not isinstance(stage_properties, dict):
        return columns
    
    # Parse TrxGenCode for column transformations
    trx_code = stage_properties.get('TrxGenCode', '')
    if trx_code:
        # Extract column assignments like: outputCol = inputCol or outputCol = expression
        patterns = [
            r'(\w+)\s*=\s*([^;\n]+)',  # Simple assignment
            r'SetOutputValue\(["\']?(\w+)["\']?,\s*([^)]+)\)',  # Function call
        ]
        for pattern in patterns:
            matches = re.findall(pattern, trx_code)
            for match in matches:
                if len(match) == 2:
                    col_name, derivation = match
                    columns[col_name.strip()] = derivation.strip()
    
    return columns

def build_variable_lineage(job_id):
    """Build field-level lineage from stages and links"""
    
    base_url = "http://localhost:8000/api"
    
    print(f"Fetching job {job_id}...")
    response = requests.get(f"{base_url}/jobs/{job_id}/full")
    if response.status_code != 200:
        print(f"Error: Job {job_id} not found")
        return []
    
    job_data = response.json()
    stages = job_data.get('stages', [])
    links = job_data.get('links', [])
    
    # Build stage lookup
    stage_map = {s['stage_id']: s for s in stages}
    stage_name_map = {s['name']: s for s in stages}
    
    # Build graph
    forward_graph = defaultdict(list)  # stage -> [(target_stage, link)]
    reverse_graph = defaultdict(list)  # stage -> [(source_stage, link)]
    
    for link in links:
        src = link.get('source_stage') or (stage_map.get(link.get('source_pin', '').split('P')[0] if link.get('source_pin') else '', {}).get('name'))
        tgt = link.get('target_stage') or (stage_map.get(link.get('target_pin', '').split('P')[0] if link.get('target_pin') else '', {}).get('name'))
        
        if src and tgt:
            forward_graph[src].append({'target': tgt, 'link': link})
            reverse_graph[tgt].append({'source': src, 'link': link})
    
    # Find output stages
    all_sources = set(forward_graph.keys())
    all_targets = set(reverse_graph.keys())
    output_stages = all_targets - all_sources
    
    print(f"Found {len(output_stages)} output stages")
    
    # Extract columns from links (field mappings)
    lineage_records = []
    
    for output_stage in output_stages:
        print(f"Analyzing: {output_stage}")
        
        # Get all fields in output stage from incoming links
        output_fields = set()
        for upstream in reverse_graph.get(output_stage, []):
            link = upstream['link']
            cols = extract_columns_from_link(link.get('properties', {}))
            for col in cols:
                if col['name']:
                    output_fields.add(col['name'])
        
        # Backtrack each field
        for field in output_fields:
            path = [output_stage]
            visited = set()
            
            def trace_field(current_stage, current_field, current_path):
                if current_stage in visited:
                    return
                visited.add(current_stage)
                
                # Get upstream connections
                upstreams = reverse_graph.get(current_stage, [])
                
                if not upstreams:
                    # Source stage
                    lineage_records.append({
                        'Target_Stage': output_stage,
                        'Target_Field': field,
                        'Source_Stage': current_stage,
                        'Source_Field': current_field,
                        'Full_Path': ' -> '.join(current_path),
                        'Total_Hops': len(current_path) - 1,
                        'Transformation': 'SOURCE'
                    })
                    return
                
                for upstream in upstreams:
                    src_stage = upstream['source']
                    link = upstream['link']
                    link_name = link.get('name', '')
                    
                    # Get column mappings from link
                    cols = extract_columns_from_link(link.get('properties', {}))
                    
                    # Check for field transformation
                    stage_cols = extract_stage_columns(stage_name_map.get(current_stage, {}).get('properties', {}))
                    
                    if current_field in stage_cols:
                        # Field has transformation
                        derivation = stage_cols[current_field]
                        
                        lineage_records.append({
                            'Target_Stage': output_stage,
                            'Target_Field': field,
                            'Source_Stage': src_stage,
                            'Source_Field': current_field,
                            'Full_Path': ' -> '.join(current_path + [src_stage]),
                            'Total_Hops': len(current_path),
                            'Transformation': derivation
                        })
                        
                        # Continue tracing
                        trace_field(src_stage, current_field, current_path + [src_stage])
                    else:
                        # Direct pass-through
                        lineage_records.append({
                            'Target_Stage': output_stage,
                            'Target_Field': field,
                            'Source_Stage': src_stage,
                            'Source_Field': current_field,
                            'Full_Path': ' -> '.join(current_path + [src_stage]),
                            'Total_Hops': len(current_path),
                            'Transformation': 'PASS_THROUGH'
                        })
                        
                        trace_field(src_stage, current_field, current_path + [src_stage])
            
            trace_field(output_stage, field, path)
    
    return lineage_records

def export_variable_lineage(job_id, output_file=None):
    """Export variable lineage to Excel"""
    
    if not output_file:
        output_file = f"job_{job_id}_variable_lineage.xlsx"
    
    lineage_data = build_variable_lineage(job_id)
    
    if not lineage_data:
        print("No variable lineage data found")
        return
    
    print(f"\nCreating Excel file: {output_file}")
    print(f"Total lineage records: {len(lineage_data)}")
    
    # Create DataFrame
    df = pd.DataFrame(lineage_data)
    
    # Sort by target stage and variable
    df = df.sort_values(['Target_Stage', 'Target_Variable', 'Depth'])
    
    # Write to Excel
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Variable Lineage', index=False)
        
        # Summary sheet
        summary = pd.DataFrame([{
            'Total Variables': df['Target_Variable'].nunique(),
            'Total Stages': df['Target_Stage'].nunique(),
            'Total Lineage Records': len(df),
            'Max Depth': df['Depth'].max() if len(df) > 0 else 0
        }])
        summary.to_excel(writer, sheet_name='Summary', index=False)
    
    # Format Excel
    wb = load_workbook(output_file)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Header formatting
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 80)
    
    wb.save(output_file)
    print(f"✓ Export complete: {output_file}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python variable_lineage.py <job_id> [output_file.xlsx]")
        sys.exit(1)
    
    job_id = int(sys.argv[1])
    output_file = sys.argv[2] if len(sys.argv) > 2 else None
    
    export_variable_lineage(job_id, output_file)
